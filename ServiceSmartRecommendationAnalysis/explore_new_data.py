import openpyxl
from datetime import datetime
from collections import Counter
import json

wb = openpyxl.load_workbook(
    "/root/ClaudeProjects/ServiceSmartRecommendationAnalysis/Ai_Report_2026-06-09(4.01-6.01).xlsx",
    data_only=True
)
ws = wb["Ai-Report"]

headers = [c.value for c in next(ws.iter_rows(max_row=1))]
col_map = {h: i for i, h in enumerate(headers)}
print("KEY COLUMNS:", {k: col_map.get(k) for k in ['createdDate', 'geo', 'configuredFrom', 'defaultFrom', 'defaultServicePn1', 'selectedServicePn1', 'country']})

dates = []
geos = Counter()
config_froms = Counter()
default_froms = Counter()
countries = Counter()

for row in ws.iter_rows(min_row=2, values_only=True):
    g = row[col_map['geo']]
    if g:
        geos[g] += 1
    
    cf = row[col_map['configuredFrom']]
    if cf:
        config_froms[cf] += 1
    
    df = row[col_map['defaultFrom']]
    if df:
        default_froms[str(df)] += 1
    
    co = row[col_map['country']]
    if co:
        countries[co] += 1
    
    date_str = row[col_map['createdDate']]
    if date_str:
        try:
            dates.append(datetime.strptime(str(date_str)[:20], '%a %b %d %H:%M:%S'))
        except:
            pass

print("\n=== DATE RANGE ===")
if dates:
    print("Min:", min(dates))
    print("Max:", max(dates))
    print("Days covered:", (max(dates)-min(dates)).days)
    
    months = Counter()
    for d in dates:
        months[d.strftime('%Y-%m')] += 1
    print("\n=== MONTHLY DISTRIBUTION ===")
    for m in sorted(months.keys()):
        print(f"  {m}: {months[m]}")

print("\n=== GEO DISTRIBUTION ===")
for g, c in geos.most_common():
    print(f"  {g}: {c}")

print("\n=== configuredFrom DISTRIBUTION ===")
for cf, c in config_froms.most_common():
    print(f"  {cf}: {c}")

print("\n=== defaultFrom DISTRIBUTION ===")
for df, c in default_froms.most_common():
    print(f"  '{df}': {c}")

print("\n=== TOP 20 COUNTRIES ===")
for co, c in countries.most_common(20):
    print(f"  {co}: {c}")

# Also check empty/defaultServicePn patterns
has_default = 0
has_selected = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    dsp = row[col_map['defaultServicePn1']]
    ssp = row[col_map['selectedServicePn1']]
    if dsp and str(dsp).strip():
        has_default += 1
    if ssp and str(ssp).strip():
        has_selected += 1

total = ws.max_row - 1
print(f"\n=== QUICK STATS ===")
print(f"Total rows: {total}")
print(f"Has defaultServicePn1: {has_default} ({has_default/total*100:.1f}%)")
print(f"Has selectedServicePn1: {has_selected} ({has_selected/total*100:.1f}%)")
