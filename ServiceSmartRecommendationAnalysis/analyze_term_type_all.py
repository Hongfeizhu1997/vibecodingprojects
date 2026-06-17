"""
Service Product Analysis by Term and Type — ALL GEOs
"""
import openpyxl
from collections import defaultdict, Counter
import re, json

FILE = "/root/ClaudeProjects/ServiceSmartRecommendationAnalysis/Ai_Report_2026-06-09(4.01-6.01).xlsx"
wb = openpyxl.load_workbook(FILE, data_only=True)
ws = wb["Ai-Report"]

headers = [c.value for c in next(ws.iter_rows(max_row=1))]
col_map = {h: i for i, h in enumerate(headers)}

TERM_PATTERNS = [
    (r'\b(\d+)Y\b', lambda m: int(m.group(1))),
    (r'\b(\d+)\s*Years?\b', lambda m: int(m.group(1))),
    (r'\b(\d+)\s*Year\b', lambda m: int(m.group(1))),
    (r'\b(\d+)YR\b', lambda m: int(m.group(1))),
    (r'\b(\d+)\s*ans?\b', lambda m: int(m.group(1))),
]

TYPE_PATTERNS = [
    ('Premier Support Plus', 'Premier Support Plus'),
    ('Premier Support', 'Premier Support'),
    ('Onsite', 'Onsite'),
    ('Basic Warranty', 'Basic Warranty'),
    ('Base Warranty', 'Base Warranty'),
    ('Warranty Extension', 'Warranty Extension'),
    ('Depot', 'Depot'),
    ('Courier', 'Courier'),
    ('Premier', 'Premier (generic)'),
    ('Essential', 'Essential'),
]

def extract_term(desc):
    if not desc: return None
    desc_str = str(desc)
    for pattern, fn in TERM_PATTERNS:
        m = re.search(pattern, desc_str, re.IGNORECASE)
        if m:
            term = fn(m)
            if 1 <= term <= 5: return term
    return None

def extract_type(desc):
    if not desc: return None
    desc_str = str(desc)
    for keyword, label in TYPE_PATTERNS:
        if keyword.lower() in desc_str.lower(): return label
    return 'Other/Unknown'

# Per-GEO stats
GEO_LIST = ['EMEA', 'AP', 'NA', 'LA']
terms_list = ['1Y','2Y','3Y','4Y','5Y']

# term_stats[geo][term] = {total, accepted, changed, no_select}
term_stats = {g: defaultdict(lambda: {'total':0,'accepted':0,'changed':0,'no_select':0}) for g in GEO_LIST}
type_stats = {g: defaultdict(lambda: {'total':0,'accepted':0,'changed':0,'no_select':0}) for g in GEO_LIST}
tt_stats = {g: defaultdict(lambda: {'total':0,'accepted':0,'changed':0,'no_select':0}) for g in GEO_LIST}
tt_acc = {g: defaultdict(lambda: defaultdict(lambda: {'total':0,'accepted':0})) for g in GEO_LIST}
trans_term = {g: defaultdict(Counter) for g in GEO_LIST}
trans_type = {g: defaultdict(Counter) for g in GEO_LIST}

for row in ws.iter_rows(min_row=2, values_only=True):
    geo = row[col_map['geo']]
    if geo not in GEO_LIST: continue
    
    dsp = row[col_map['defaultServicePn1']]
    dsp_desc = row[col_map['defaultServicePn1Desc']]
    ssp = row[col_map['selectedServicePn1']]
    ssp_desc = row[col_map['selectedServicePn1Desc']]
    
    has_default = bool(dsp and str(dsp).strip())
    has_selected = bool(ssp and str(ssp).strip())
    if not has_default: continue
    
    d_term = extract_term(dsp_desc)
    d_type = extract_type(dsp_desc)
    s_term = extract_term(ssp_desc) if has_selected else None
    s_type = extract_type(ssp_desc) if has_selected else None
    
    accepted = has_default and has_selected and (str(dsp).strip() == str(ssp).strip())
    changed = has_default and has_selected and (str(dsp).strip() != str(ssp).strip())
    no_select = has_default and not has_selected
    
    if d_term:
        tkey = f"{d_term}Y"
        term_stats[geo][tkey]['total'] += 1
        if accepted: term_stats[geo][tkey]['accepted'] += 1
        if changed: term_stats[geo][tkey]['changed'] += 1
        if no_select: term_stats[geo][tkey]['no_select'] += 1
        if s_term: trans_term[geo][tkey][f"{s_term}Y"] += 1
    
    if d_type:
        type_stats[geo][d_type]['total'] += 1
        if accepted: type_stats[geo][d_type]['accepted'] += 1
        if changed: type_stats[geo][d_type]['changed'] += 1
        if no_select: type_stats[geo][d_type]['no_select'] += 1
        if s_type: trans_type[geo][d_type][s_type] += 1
    
    if d_term and d_type:
        tkey = f"{d_term}Y"
        ttkey = f"{tkey} {d_type}"
        tt_stats[geo][ttkey]['total'] += 1
        if accepted: tt_stats[geo][ttkey]['accepted'] += 1
        if changed: tt_stats[geo][ttkey]['changed'] += 1
        if no_select: tt_stats[geo][ttkey]['no_select'] += 1
        tt_acc[geo][tkey][d_type]['total'] += 1
        if accepted: tt_acc[geo][tkey][d_type]['accepted'] += 1

def pct(num, den):
    return f"{num/den*100:.1f}%" if den else "—"

# ==================== PRINT ====================
for geo in GEO_LIST:
    ts = term_stats[geo]
    tys = type_stats[geo]
    
    print(f"\n{'='*80}")
    print(f"{geo} — TERM ANALYSIS")
    print(f"{'='*80}")
    print(f"{'Term':<8} {'Count':>7} {'Accept':>7} {'Acc%':>7} {'Changed':>7} {'Chg%':>7} {'NoSel':>7} {'NS%':>7}")
    print("-"*70)
    for t in terms_list:
        s = ts[t]; ttl = s['total']
        if ttl == 0: continue
        print(f"{t:<8} {ttl:>7} {s['accepted']:>7} {pct(s['accepted'],ttl):>7} {s['changed']:>7} {pct(s['changed'],ttl):>7} {s['no_select']:>7} {pct(s['no_select'],ttl):>7}")
    
    print(f"\n{geo} — TYPE ANALYSIS")
    print(f"{'='*80}")
    print(f"{'Type':<25} {'Count':>7} {'Accept':>7} {'Acc%':>7} {'Changed':>7} {'Chg%':>7} {'NoSel':>7} {'NS%':>7}")
    print("-"*70)
    for tp in sorted(tys.keys(), key=lambda k: tys[k]['total'], reverse=True):
        s = tys[tp]; ttl = s['total']
        if ttl < 20: continue
        print(f"{tp:<25} {ttl:>7} {s['accepted']:>7} {pct(s['accepted'],ttl):>7} {s['changed']:>7} {pct(s['changed'],ttl):>7} {s['no_select']:>7} {pct(s['no_select'],ttl):>7}")
    
    print(f"\n{geo} — TERM × TYPE MATRIX")
    print(f"{'='*80}")
    types_sorted = sorted(tys.keys(), key=lambda k: tys[k]['total'], reverse=True)
    active_types = [tp for tp in types_sorted if tys[tp]['total'] >= 30]
    active_terms = [t for t in terms_list if ts[t]['total'] >= 20]
    
    hdr = "        "
    for tp in active_types: hdr += f"{tp:>18}"
    print(hdr)
    for t in active_terms:
        row_str = f"{t:<8}"
        for tp in active_types:
            d = tt_acc[geo][t][tp]
            if d['total'] >= 10:
                acc = d['accepted']/d['total']*100
                row_str += f"{acc:>6.1f}%({d['total']:>4})  "
            else:
                row_str += f"{'':>18}"
        print(row_str)
    
    # Top combos
    print(f"\n{geo} — TOP 10 TERM × TYPE COMBOS")
    print(f"{'='*80}")
    print(f"{'Combination':<40} {'Count':>7} {'Acc%':>8} {'Chg%':>8}")
    print("-"*60)
    for combo in sorted(tt_stats[geo].keys(), key=lambda k: tt_stats[geo][k]['total'], reverse=True)[:10]:
        s = tt_stats[geo][combo]; ttl = s['total']
        print(f"{combo:<40} {ttl:>7} {pct(s['accepted'],ttl):>8} {pct(s['changed'],ttl):>8}")

# Cross-GEO summary for key combos
print(f"\n{'='*80}")
print(f"CROSS-GEO COMPARISON — 3Y Premier Support")
print(f"{'='*80}")
for geo in GEO_LIST:
    s = tt_stats[geo].get('3Y Premier Support', {'total':0,'accepted':0,'changed':0})
    ttl = s['total']
    print(f"  {geo}: {ttl:>6} configs | Acc={pct(s['accepted'],ttl):>6} | Chg={pct(s['changed'],ttl):>6}")

print(f"\nCROSS-GEO — 3Y Onsite")
for geo in GEO_LIST:
    s = tt_stats[geo].get('3Y Onsite', {'total':0,'accepted':0,'changed':0})
    ttl = s['total']
    print(f"  {geo}: {ttl:>6} configs | Acc={pct(s['accepted'],ttl):>6} | Chg={pct(s['changed'],ttl):>6}")

print(f"\nCROSS-GEO — Premier Support Plus (all terms)")
for geo in GEO_LIST:
    s = type_stats[geo].get('Premier Support Plus', {'total':0,'accepted':0,'changed':0})
    ttl = s['total']
    print(f"  {geo}: {ttl:>6} configs | Acc={pct(s['accepted'],ttl):>6} | Chg={pct(s['changed'],ttl):>6}")
