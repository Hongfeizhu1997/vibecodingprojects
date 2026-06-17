import openpyxl

wb = openpyxl.load_workbook(
    "/root/ClaudeProjects/ServiceSmartRecommendationAnalysis/Ai_Report_2026-06-09(4.01-6.01).xlsx",
    data_only=True
)
ws = wb["Ai-Report"]

headers = [c.value for c in next(ws.iter_rows(max_row=1))]
col_map = {h: i for i, h in enumerate(headers)}

# Print raw date samples
print("=== RAW createdDate SAMPLES ===")
count = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    d = row[col_map['createdDate']]
    if d and count < 10:
        print(f"  type={type(d).__name__}, value='{d}'")
        count += 1

# Also check modifiedDate
print("\n=== RAW modifiedDate SAMPLES ===")
count = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    d = row[col_map['modifiedDate']]
    if d and count < 10:
        print(f"  type={type(d).__name__}, value='{d}'")
        count += 1
