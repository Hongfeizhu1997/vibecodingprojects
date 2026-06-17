import openpyxl
from datetime import datetime
from collections import Counter

files = [
    ("Ai_Report_20250815~20260306.xlsx", "Aug 15,2025 - Mar 6,2026"),
    ("Ai_Report_2026-04-01(3.17-4.01).xlsx", "Mar 17 - Apr 1,2026"),
]

for fname, label in files:
    try:
        wb = openpyxl.load_workbook(
            f"/root/ClaudeProjects/ServiceSmartRecommendationAnalysis/{fname}",
            data_only=True
        )
        print(f"\n=== {label} ===")
        print(f"Sheets: {wb.sheetnames}")
        for sn in wb.sheetnames:
            ws = wb[sn]
            print(f"  {sn}: {ws.max_row} rows x {ws.max_column} cols")
    except Exception as e:
        print(f"\n=== {label} === ERROR: {e}")
