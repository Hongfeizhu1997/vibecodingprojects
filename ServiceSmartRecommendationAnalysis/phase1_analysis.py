"""
Phase 1 Operational KPI Analysis — New Data (Apr 1 - Jun 1, 2026)
Computes all metrics across all dimensions.
"""
import openpyxl
from datetime import datetime
from collections import defaultdict, Counter
import json

# =============================================================================
# 1. LOAD NEW DATA
# =============================================================================
NEW_FILE = "/root/ClaudeProjects/ServiceSmartRecommendationAnalysis/Ai_Report_2026-06-09(4.01-6.01).xlsx"
wb = openpyxl.load_workbook(NEW_FILE, data_only=True)
ws = wb["Ai-Report"]

headers = [c.value for c in next(ws.iter_rows(max_row=1))]
col_map = {h: i for i, h in enumerate(headers)}

print("Columns:", list(col_map.keys()))

# =============================================================================
# 2. PARSE ALL ROWS
# =============================================================================
metrics = {}  # key = composite dim key -> counts

def init_counts():
    return {
        'total': 0,
        'has_default': 0,         # Coverage numerator
        'ai_default': 0,           # AI Coverage numerator
        'has_selected': 0,         # Attach numerator
        'accepted': 0,             # selected == default
        'changed': 0,              # selected != default (both present)
        'attach_failure': 0,       # has default but no selected
        'new_cto': 0,
        'copied_cto': 0,
        'mtm': 0,
    }

overall = init_counts()
geo_data = defaultdict(init_counts)
month_data = defaultdict(init_counts)
geo_month_data = defaultdict(lambda: defaultdict(init_counts))
cf_data = defaultdict(init_counts)      # configuredFrom
df_data = defaultdict(init_counts)      # defaultFrom
geo_cf_data = defaultdict(lambda: defaultdict(init_counts))

for row in ws.iter_rows(min_row=2, values_only=True):
    geo = row[col_map['geo']]
    cf = row[col_map['configuredFrom']]  # New CTO / Copied CTO / MTM
    df = str(row[col_map['defaultFrom']]) if row[col_map['defaultFrom']] else ''
    dsp = row[col_map['defaultServicePn1']]
    ssp = row[col_map['selectedServicePn1']]
    date_str = row[col_map['createdDate']]
    
    # Parse date
    month = None
    if date_str:
        try:
            dt = datetime.strptime(str(date_str).strip(), '%a %b %d %H:%M:%S UTC %Y')
            month = dt.strftime('%Y-%m')
        except:
            pass
    
    if not geo:
        continue
    
    has_default = bool(dsp and str(dsp).strip())
    has_selected = bool(ssp and str(ssp).strip())
    is_ai = (df == 'AI')
    accepted = has_default and has_selected and (str(dsp).strip() == str(ssp).strip())
    changed = has_default and has_selected and (str(dsp).strip() != str(ssp).strip())
    attach_fail = has_default and not has_selected
    
    # Update overall
    overall['total'] += 1
    if has_default: overall['has_default'] += 1
    if is_ai: overall['ai_default'] += 1
    if has_selected: overall['has_selected'] += 1
    if accepted: overall['accepted'] += 1
    if changed: overall['changed'] += 1
    if attach_fail: overall['attach_failure'] += 1
    if cf == 'New CTO': overall['new_cto'] += 1
    elif cf == 'Copied CTO': overall['copied_cto'] += 1
    elif cf == 'MTM': overall['mtm'] += 1
    
    # Update geo
    gd = geo_data[geo]
    gd['total'] += 1
    if has_default: gd['has_default'] += 1
    if is_ai: gd['ai_default'] += 1
    if has_selected: gd['has_selected'] += 1
    if accepted: gd['accepted'] += 1
    if changed: gd['changed'] += 1
    if attach_fail: gd['attach_failure'] += 1
    if cf == 'New CTO': gd['new_cto'] += 1
    elif cf == 'Copied CTO': gd['copied_cto'] += 1
    elif cf == 'MTM': gd['mtm'] += 1
    
    # Update month
    if month:
        md = month_data[month]
        md['total'] += 1
        if has_default: md['has_default'] += 1
        if is_ai: md['ai_default'] += 1
        if has_selected: md['has_selected'] += 1
        if accepted: md['accepted'] += 1
        if changed: md['changed'] += 1
        if attach_fail: md['attach_failure'] += 1
        
        # geo-month
        gmd = geo_month_data[geo][month]
        gmd['total'] += 1
        if has_default: gmd['has_default'] += 1
        if is_ai: gmd['ai_default'] += 1
        if has_selected: gmd['has_selected'] += 1
        if accepted: gmd['accepted'] += 1
        if changed: gmd['changed'] += 1
        if attach_fail: gmd['attach_failure'] += 1
    
    # Update configuredFrom
    if cf:
        cfd = cf_data[cf]
        cfd['total'] += 1
        if has_default: cfd['has_default'] += 1
        if is_ai: cfd['ai_default'] += 1
        if has_selected: cfd['has_selected'] += 1
        if accepted: cfd['accepted'] += 1
        if changed: cfd['changed'] += 1
        if attach_fail: cfd['attach_failure'] += 1
    
    # Update defaultFrom
    if df:
        dfd = df_data[df]
        dfd['total'] += 1
        if has_default: dfd['has_default'] += 1
        if is_ai: dfd['ai_default'] += 1
        if has_selected: dfd['has_selected'] += 1
        if accepted: dfd['accepted'] += 1
        if changed: dfd['changed'] += 1
        if attach_fail: dfd['attach_failure'] += 1
    
    # GEO x configuredFrom
    if geo and cf:
        gcf = geo_cf_data[geo][cf]
        gcf['total'] += 1
        if has_default: gcf['has_default'] += 1
        if has_selected: gcf['has_selected'] += 1
        if accepted: gcf['accepted'] += 1

# =============================================================================
# 3. COMPUTE RATES
# =============================================================================
def compute_rates(d):
    t = d['total']
    hd = d['has_default']
    return {
        'total': t,
        'coverage': round(hd / t * 100, 1) if t else 0,
        'ai_coverage': round(d['ai_default'] / t * 100, 1) if t else 0,
        'attach': round(d['has_selected'] / t * 100, 1) if t else 0,
        'acceptance': round(d['accepted'] / hd * 100, 1) if hd else 0,
        'change': round(d['changed'] / hd * 100, 1) if hd else 0,
        'attach_failure': round(d['attach_failure'] / hd * 100, 1) if hd else 0,
        'has_default': hd,
        'has_selected': d['has_selected'],
        'accepted': d['accepted'],
        'changed': d['changed'],
        'new_cto': d['new_cto'],
        'copied_cto': d['copied_cto'],
        'mtm': d['mtm'],
    }

print("\n" + "="*80)
print("PHASE 1: OVERALL METRICS (Apr 1 - Jun 1, 2026)")
print("="*80)
ov = compute_rates(overall)
for k, v in ov.items():
    print(f"  {k}: {v}")

print("\n" + "="*80)
print("BY GEO")
print("="*80)
geo_rates = {}
for geo in sorted(geo_data.keys()):
    gr = compute_rates(geo_data[geo])
    geo_rates[geo] = gr
    print(f"\n  --- {geo} (n={gr['total']}) ---")
    print(f"  Coverage:     {gr['coverage']}%")
    print(f"  AI Coverage:  {gr['ai_coverage']}%")
    print(f"  Attach Rate:  {gr['attach']}%")
    print(f"  Acceptance:   {gr['acceptance']}%")
    print(f"  Change Rate:  {gr['change']}%")
    print(f"  Attach Fail:  {gr['attach_failure']}%")
    print(f"  New CTO: {gr['new_cto']}, Copied: {gr['copied_cto']}, MTM: {gr['mtm']}")

print("\n" + "="*80)
print("BY MONTH")
print("="*80)
month_rates = {}
for m in sorted(month_data.keys()):
    mr = compute_rates(month_data[m])
    month_rates[m] = mr
    print(f"\n  --- {m} (n={mr['total']}) ---")
    for k in ['coverage', 'ai_coverage', 'attach', 'acceptance', 'change', 'attach_failure']:
        print(f"  {k}: {mr[k]}%")

print("\n" + "="*80)
print("BY configuredFrom")
print("="*80)
for cf in sorted(cf_data.keys()):
    cfr = compute_rates(cf_data[cf])
    print(f"\n  --- {cf} (n={cfr['total']}) ---")
    for k in ['coverage', 'ai_coverage', 'attach', 'acceptance', 'change']:
        print(f"  {k}: {cfr[k]}%")

print("\n" + "="*80)
print("BY defaultFrom")
print("="*80)
for df in sorted(df_data.keys()):
    dfr = compute_rates(df_data[df])
    print(f"\n  --- {df} (n={dfr['total']}) ---")
    for k in ['coverage', 'attach', 'acceptance', 'change']:
        print(f"  {k}: {dfr[k]}%")

print("\n" + "="*80)
print("GEO × MONTH — Acceptance Rate")
print("="*80)
for geo in sorted(geo_month_data.keys()):
    print(f"\n  {geo}:")
    for m in sorted(geo_month_data[geo].keys()):
        gmr = compute_rates(geo_month_data[geo][m])
        print(f"    {m}: Acceptance={gmr['acceptance']}%, Attach={gmr['attach']}%, Coverage={gmr['coverage']}% (n={gmr['total']})")

print("\n" + "="*80)
print("GEO × configuredFrom — Acceptance Rate")
print("="*80)
for geo in sorted(geo_cf_data.keys()):
    print(f"\n  {geo}:")
    for cf in sorted(geo_cf_data[geo].keys()):
        gcfr = compute_rates(geo_cf_data[geo][cf])
        print(f"    {cf}: Acceptance={gcfr['acceptance']}%, n={gcfr['total']}")

# Save all data for HTML report
output = {
    'overall': ov,
    'geo': {g: geo_rates[g] for g in geo_rates},
    'month': {m: month_rates[m] for m in month_rates},
}
with open('/tmp/phase1_new_data.json', 'w') as f:
    json.dump(output, f, indent=2, default=str)
print("\nData saved to /tmp/phase1_new_data.json")
