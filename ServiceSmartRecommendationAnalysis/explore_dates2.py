import openpyxl
from datetime import datetime
from collections import Counter

wb = openpyxl.load_workbook(
    "/root/ClaudeProjects/ServiceSmartRecommendationAnalysis/Ai_Report_2026-06-09(4.01-6.01).xlsx",
    data_only=True
)
ws = wb["Ai-Report"]
headers = [c.value for c in next(ws.iter_rows(max_row=1))]
col_map = {h: i for i, h in enumerate(headers)}

dates = []
for row in ws.iter_rows(min_row=2, values_only=True):
    d = row[col_map['createdDate']]
    if d:
        try:
            dates.append(datetime.strptime(str(d).strip(), '%a %b %d %H:%M:%S UTC %Y'))
        except:
            pass

print("Min:", min(dates))
print("Max:", max(dates))
print("Days:", (max(dates)-min(dates)).days)

months = Counter()
for d in dates:
    months[d.strftime('%Y-%m')] += 1
print("\nMonthly distribution:")
for m in sorted(months.keys()):
    print(f"  {m}: {months[m]}")
