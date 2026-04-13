"""
Generate sample_pfv.xlsx — a PFV matrix file for Stage 5 import.
Sheet name: "Coparison" (matches the original file's typo so Stage 5 finds it).

Column layout (0-based index → Stage 5 field):
  A(0) = tier label / blank
  B(1) = YR (year 1-5)
  C(2) = kyd
  D(3) = ti cru  (ticru)
  E(4) = adp
  F(5) = sbty 3yr (sbty)
  G(6) = 1YR cciBW depot       (d1c_depot)
  H(7) = 1YR cciBW exp depot   (d1c_exp)
  I(8) = 1YR cciBW onsite      (d1c_onsite)
  J(9) = 3YR cciBW depot       (d3c_depot)
  K(10)= 3YR cciBW exp depot   (d3c_exp)
  L(11)= 3YR cciBW onsite      (d3c_onsite)
  M(12)= 3YR OS BW onsite      (d3o_onsite)
  N(13)= 4YR cciBW onsite      (d4c_onsite)
  O(14)= 4YR OS BW onsite      (d4o_onsite)
  P(15)= Priority               (priority)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# PFV DATA — (kyd, ticru, adp, sbty, d1c_dep, d1c_exp, d1c_on,
#              d3c_dep, d3c_exp, d3c_on, d3o_on, d4c_on, d4o_on, priority)
# Year rows 1-5 for each tier
# ---------------------------------------------------------------------------
TIERS = {
    'Tablet': [
        # yr  kyd  ticru  adp  sbty  d1c_dep d1c_exp d1c_on  d3c_dep d3c_exp d3c_on  d3o_on  d4c_on  d4o_on  priority
        (1,    0,    0,   40,    0,      0,     20,     0,       0,      0,      0,      0,      0,      0,      75),
        (2,    0,    0,   70,   25,     55,     70,     0,       0,      0,      0,      0,      0,      0,      85),
        (3,    0,    0,  100,   25,    120,    140,     0,       0,      0,      0,      0,      0,      0,      95),
        (4,    0,    0,  130,   25,    185,    210,     0,       0,      0,      0,      0,      0,      0,     105),
        (5,    0,    0,  160,   25,    250,    280,     0,       0,      0,      0,      0,      0,      0,     115),
    ],
    'Premium': [
        (1,    9,   10,   40,    0,      0,     20,    25,       0,      0,      0,      0,      0,      0,      70),
        (2,   15,   20,   65,    0,     45,     60,    80,       0,      0,      0,      0,      0,      0,      80),
        (3,   27,   35,   90,   25,     95,    125,   150,       0,     40,     60,      0,      0,      0,      90),
        (4,   35,   50,  120,   25,    150,    185,   220,      60,    115,    130,     80,     70,      0,     100),
        (5,   45,   65,  155,   25,    220,    255,   285,     130,    180,    230,    165,    160,     95,     110),
    ],
    'Classic': [
        (1,    9,   10,   30,    0,      0,     10,    15,       0,      0,      0,      0,      0,      0,      70),
        (2,   15,   20,   60,    0,     35,     40,    50,       0,      0,      0,      0,      0,      0,      80),
        (3,   27,   35,   90,   20,     75,     80,   100,       0,     25,     45,      0,      0,      0,      90),
        (4,   35,   50,  120,   20,    135,    140,   150,      40,     75,     90,     60,     70,      0,     100),
        (5,   45,   65,  155,   20,    205,    210,   220,      90,    125,    160,    120,    120,     70,     110),
    ],
    'Edge': [
        (1,    7,    8,   25,    0,      0,      8,    12,       0,      0,      0,      0,      0,      0,      60),
        (2,   12,   16,   50,    0,     28,     33,    42,       0,      0,      0,      0,      0,      0,      70),
        (3,   22,   28,   75,   15,     62,     67,    82,       0,     20,     37,      0,      0,      0,      80),
        (4,   28,   40,  100,   15,    110,    115,   125,      32,     60,     74,     48,     55,      0,      90),
        (5,   36,   52,  130,   15,    168,    173,   183,      73,    102,    132,     96,     97,     55,     100),
    ],
    'Education': [
        (1,    5,    6,   20,    0,      0,      6,     9,       0,      0,      0,      0,      0,      0,      55),
        (2,    9,   12,   40,    0,     22,     26,    34,       0,      0,      0,      0,      0,      0,      62),
        (3,   17,   22,   60,   12,     50,     54,    66,       0,     16,     30,      0,      0,      0,      70),
        (4,   22,   32,   80,   12,     88,     92,   100,      25,     48,     60,     38,     44,      0,      78),
        (5,   29,   42,  104,   12,    135,    138,   146,      58,     82,    106,     77,     78,     44,      88),
    ],
}

# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------
TIER_FILL   = PatternFill('solid', fgColor='1F4E79')  # dark blue
TIER_FONT   = Font(bold=True, color='FFFFFF', size=11)
CAT_FILL    = PatternFill('solid', fgColor='2E75B6')  # medium blue
CAT_FONT    = Font(bold=True, color='FFFFFF', size=9)
HDR_FILL    = PatternFill('solid', fgColor='D6E4F0')  # light blue
HDR_FONT    = Font(bold=True, size=9)
DATA_FONT   = Font(size=9)
NUM_FILL_ODD  = PatternFill('solid', fgColor='EBF3FB')
NUM_FILL_EVEN = PatternFill('solid', fgColor='FFFFFF')

thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CENTER = Alignment(horizontal='center', vertical='center')
LEFT   = Alignment(horizontal='left',   vertical='center')

# ---------------------------------------------------------------------------
# Build workbook
# ---------------------------------------------------------------------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Coparison'   # intentional typo matching original file

# Set column widths
ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 5
for col_letter in ['C','D','E','F']:
    ws.column_dimensions[col_letter].width = 8
for col_letter in ['G','H','I','J','K','L','M','N','O','P']:
    ws.column_dimensions[col_letter].width = 9

def write_tier_block(ws, start_row, tier_name, rows_data):
    """Write one tier block starting at start_row. Returns next free row."""
    r = start_row

    # --- Row 1: Tier name header ---
    ws.cell(r, 1, tier_name).font   = TIER_FONT
    ws.cell(r, 1).fill              = TIER_FILL
    ws.cell(r, 1).alignment         = LEFT
    ws.cell(r, 1).border            = BORDER
    for c in range(2, 17):
        ws.cell(r, c).fill   = TIER_FILL
        ws.cell(r, c).border = BORDER
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=16)
    r += 1

    # --- Row 2: Category span headers ---
    cat_headers = {
        1:  '',
        2:  'YR',
        3:  'Add-ons',
        7:  '1YR cciBW',
        10: '3YR cciBW',
        13: '3YR OS BW',
        14: '4YR cciBW',
        15: '4YR OS BW',
        16: 'Priority',
    }
    spans = [
        (1, 1),    # A
        (2, 2),    # B
        (3, 6),    # C-F  Add-ons
        (7, 9),    # G-I  1YR cciBW
        (10, 12),  # J-L  3YR cciBW
        (13, 13),  # M    3YR OS BW
        (14, 14),  # N    4YR cciBW
        (15, 15),  # O    4YR OS BW
        (16, 16),  # P    Priority
    ]
    labels = ['', 'YR', 'Add-ons', '1YR cciBW', '3YR cciBW', '3YR OS BW', '4YR cciBW', '4YR OS BW', 'Priority']
    for (sc, ec), label in zip(spans, labels):
        ws.cell(r, sc, label).font      = CAT_FONT
        ws.cell(r, sc).fill             = CAT_FILL
        ws.cell(r, sc).alignment        = CENTER
        ws.cell(r, sc).border           = BORDER
        for c in range(sc, ec + 1):
            ws.cell(r, c).fill   = CAT_FILL
            ws.cell(r, c).border = BORDER
        if ec > sc:
            ws.merge_cells(start_row=r, start_column=sc, end_row=r, end_column=ec)
    r += 1

    # --- Row 3: Detailed column headers ---
    col_headers = ['', 'YR', 'kyd', 'ti cru', 'adp', 'sbty 3yr',
                   'depot', 'exp depot', 'onsite',
                   'depot', 'exp depot', 'onsite',
                   'onsite', 'onsite', 'onsite', 'priority']
    for ci, hdr in enumerate(col_headers, start=1):
        cell = ws.cell(r, ci, hdr)
        cell.font      = HDR_FONT
        cell.fill      = HDR_FILL
        cell.alignment = CENTER
        cell.border    = BORDER
    r += 1

    # --- Data rows (years 1-5) ---
    for i, yr_row in enumerate(rows_data):
        fill = NUM_FILL_ODD if i % 2 == 0 else NUM_FILL_EVEN
        yr = yr_row[0]
        values = [None, yr] + list(yr_row[1:])  # col A blank, col B = year
        for ci, val in enumerate(values, start=1):
            cell = ws.cell(r, ci, val)
            cell.font      = DATA_FONT
            cell.fill      = fill
            cell.alignment = CENTER
            cell.border    = BORDER
        r += 1

    # --- Blank separator row ---
    r += 1
    return r


current_row = 1
for tier_name, data in TIERS.items():
    current_row = write_tier_block(ws, current_row, tier_name, data)

# Freeze top row (none — the tier blocks shift around, no fixed header to freeze)
ws.freeze_panes = None

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
out_path = r'C:\Users\Administrator\ClaudeProjects\GPS_Cost_Pricing_Workbench\stage5_browser\sample_pfv.xlsx'
wb.save(out_path)
print(f'Saved: {out_path}')
print(f'Tiers written: {list(TIERS.keys())}')
print(f'Rows per tier: 1 (name) + 1 (category) + 1 (header) + 5 (data) + 1 (blank) = 9 rows each')
