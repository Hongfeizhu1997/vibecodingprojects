#!/usr/bin/env python3
"""Extract text + tables from the first 30 slides of the blueprint PPT."""
from pptx import Presentation
from pptx.util import Emu

PPT = "/root/ClaudeProjects/SSG_AI/SSG_Sales_AI_Blueprint/references/SSG_Sales_AI_Blueprint_v4_edited.pptx"

prs = Presentation(PPT)
print(f"TOTAL SLIDES: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides):
    if i >= 30:
        break
    print(f"\n{'='*70}")
    print(f"===== SLIDE {i+1} =====")
    print(f"{'='*70}")
    # shapes in order
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = []
            for para in shape.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    txt.append(t)
            if txt:
                print("\n".join(txt))
        if getattr(shape, "has_table", False) and shape.has_table:
            print("[TABLE]")
            for row in shape.table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                print(" | ".join(cells))
    # notes
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            print(f"[NOTES]: {notes}")
